from fastapi import FastAPI, Request
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import List, Optional, Dict, Any
from datetime import datetime
import os
import json
import pandas as pd
from openpyxl import Workbook, load_workbook
import random
from dotenv import load_dotenv
import re
from fastapi.responses import StreamingResponse
load_dotenv()

app = FastAPI()

# Allow CORS for local frontend dev
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- Placeholder Feedback Questions (RAG) ---
FEEDBACK_QUESTIONS = [
    "How would you rate your overall experience at LEGOLAND Discovery Center Toronto?",
    "Which attraction or activity did you enjoy the most?",
    "What could we improve to make your visit even better?",
    "Would you recommend LEGOLAND Discovery Center Toronto to friends and family? Why or why not?",
    "Any additional comments about your visit or suggestions for future improvements?"
]

# --- LEGOLAND Discovery Center Toronto Company Information (RAG) ---
COMPANY_DOCS = [
    "LEGOLAND Discovery Center Toronto is located at Vaughan Mills Shopping Centre in Vaughan, Ontario, Canada.",
    "Our center features over 10 interactive attractions including MINILAND, LEGO Factory Tour, and Kingdom Quest laser ride.",
    "We offer birthday party packages and educational programs for schools and groups.",
    "Children under 2 years old receive free admission to LEGOLAND Discovery Center Toronto.",
    "Our center is designed for families with children aged 3-10, but all ages are welcome to enjoy the experience.",
    "We have a LEGO retail store where visitors can purchase LEGO sets, exclusive merchandise, and souvenirs.",
    "LEGOLAND Discovery Center Toronto offers annual passes for frequent visitors with special member benefits.",
    "Our center is wheelchair accessible and we provide accommodations for guests with special needs.",
    "We host seasonal events and activities throughout the year including holiday celebrations and themed weekends.",
    "Food and beverages are available at our café, and outside food is not permitted in the center.",
    "Photography is allowed throughout the center, but flash photography is restricted in certain areas.",
    "We offer online ticket booking with time slots to manage capacity and ensure a great experience for all guests.",
    "Our staff are trained to provide excellent customer service and ensure guest safety at all attractions.",
    "LEGOLAND Discovery Center Toronto is part of the global LEGOLAND family with locations worldwide.",
    "We are committed to providing a safe, clean, and fun environment for families to create lasting memories together."
]

# --- Models ---
class ChatRequest(BaseModel):
    messages: List[dict]  # [{role: 'user'|'assistant', content: str}]

class FeedbackRequest(BaseModel):
    answers: Dict[str, Any]  # {question: answer or {answer, remark}}
    conversation: Optional[List[dict]] = None

class RAGRequest(BaseModel):
    query: str

# --- Helper Functions ---
FEEDBACK_XLSX = "feedback.xlsx"
FEEDBACK_CSV = "feedback.csv"
FEEDBACK_JSON = "feedback.json"

def save_structured_feedback(answers: Dict[str, str], conversation: Optional[List[dict]]):
    timestamp = datetime.utcnow().isoformat()
    entry = {"timestamp": timestamp, **answers, "conversation": json.dumps(conversation or [])}
    # Save to Excel
    if not os.path.exists(FEEDBACK_XLSX):
        wb = Workbook()
        ws = wb.active
        headers = ["timestamp"] + list(answers.keys()) + ["conversation"]
        ws.append(headers)
        wb.save(FEEDBACK_XLSX)
    wb = load_workbook(FEEDBACK_XLSX)
    ws = wb.active
    # Ensure all columns exist
    existing_headers = [cell.value for cell in ws[1]]
    for key in answers.keys():
        if key not in existing_headers:
            ws.cell(row=1, column=len(existing_headers)+1, value=key)
            existing_headers.append(key)
    # Prepare row in correct order
    row = [timestamp] + [answers.get(h, "") for h in existing_headers if h not in ("timestamp", "conversation")] + [json.dumps(conversation or [])]
    ws.append(row)
    wb.save(FEEDBACK_XLSX)
    # Save to CSV
    df = pd.DataFrame([entry])
    if not os.path.exists(FEEDBACK_CSV):
        df.to_csv(FEEDBACK_CSV, index=False, mode='w')
    else:
        df.to_csv(FEEDBACK_CSV, index=False, mode='a', header=False)
    # Save to JSON (append)
    if not os.path.exists(FEEDBACK_JSON):
        with open(FEEDBACK_JSON, 'w') as f:
            json.dump([entry], f, indent=2)
    else:
        with open(FEEDBACK_JSON, 'r+') as f:
            try:
                data = json.load(f)
            except json.JSONDecodeError:
                data = []
            data.append(entry)
            f.seek(0)
            json.dump(data, f, indent=2)
            f.truncate()

# --- Endpoints ---

@app.get("/health")
def health_check():
    """Health check endpoint."""
    return {"status": "ok"}

@app.post("/chat")
def chat_endpoint(req: ChatRequest):
    """Chat endpoint for processing user messages with streaming."""
    try:
        from together import Together
    except ImportError:
        def error_stream():
            yield "[Error: together SDK not installed on backend. Please install with: pip install together]"
        return StreamingResponse(error_stream(), media_type="text/plain")
    TOGETHER_API_KEY = os.getenv("TOGETHER_API_KEY")
    if not TOGETHER_API_KEY:
        def error_stream():
            yield "[Error: TOGETHER_API_KEY not set in environment]"
        return StreamingResponse(error_stream(), media_type="text/plain")
    client = Together(api_key=TOGETHER_API_KEY)
    feedback_schema = {
        "type": "object",
        "properties": {
            "conversation": {"type": "string", "description": "Conversational message for the user."},
            "data": {
                "type": "object",
                "properties": {
                    "Q1": {"type": "object", "properties": {"answer": {"type": "string"}, "remark": {"type": "string"}}, "required": ["answer", "remark"]},
                    "Q2": {"type": "object", "properties": {"answer": {"type": "string"}, "remark": {"type": "string"}}, "required": ["answer", "remark"]},
                    "Q3": {"type": "object", "properties": {"answer": {"type": "string"}, "remark": {"type": "string"}}, "required": ["answer", "remark"]},
                    "Q4": {"type": "object", "properties": {"answer": {"type": "string"}, "remark": {"type": "string"}}, "required": ["answer", "remark"]},
                    "Q5": {"type": "object", "properties": {"answer": {"type": "string"}, "remark": {"type": "string"}}, "required": ["answer", "remark"]}
                },
                "required": ["Q1", "Q2", "Q3", "Q4", "Q5"]
            },
            "current_question": {"type": "string"},
            "is_complete": {"type": "boolean"}
        },
        "required": ["conversation", "data", "current_question", "is_complete"]
    }
    def together_stream():
        try:
            response = client.chat.completions.create(
                model="meta-llama/Llama-3.3-70B-Instruct-Turbo-Free",
                messages=req.messages,
                response_format={
                    "type": "json_object",
                    "schema": feedback_schema
                },
                stream=True
            )
            for chunk in response:
                # Each chunk is a ChatCompletionStreamOutput
                if hasattr(chunk, "choices") and chunk.choices:
                    delta = chunk.choices[0].delta
                    if hasattr(delta, "content") and delta.content:
                        yield delta.content
        except Exception as e:
            yield f"[Error from Together API: {str(e)}]"
    return StreamingResponse(together_stream(), media_type="text/plain")

@app.post("/feedback")
def feedback_endpoint(req: FeedbackRequest):
    """Feedback saving endpoint (structured, extracts answers and remarks from conversation data)."""
    # Extract answers and remarks from the conversation data structure
    answers = {}
    remarks = {}
    if req.answers:
        for i, q in enumerate(FEEDBACK_QUESTIONS, 1):
            q_key = f"Q{i}"
            if q_key in req.answers:
                q_data = req.answers[q_key]
                if isinstance(q_data, dict):
                    answers[q] = q_data.get("answer", "")
                    remarks[f"{q} - Remark"] = q_data.get("remark", "")
                else:
                    answers[q] = str(q_data)
                    remarks[f"{q} - Remark"] = ""
            else:
                answers[q] = ""
                remarks[f"{q} - Remark"] = ""
    
    # Combine answers and remarks for saving
    all_data = {**answers, **remarks}
    save_structured_feedback(all_data, req.conversation)
    return {"status": "saved", "answers": answers, "remarks": remarks}

@app.post("/rag")
def rag_endpoint(req: RAGRequest):
    """RAG endpoint returns feedback questions or relevant doc/FAQ."""
    query = req.query.lower().strip()
    # If the query is 'feedback', return the questions
    if query == "feedback":
        return {"questions": FEEDBACK_QUESTIONS}
    # Otherwise, return the most relevant doc/FAQ (simple keyword match)
    for doc in COMPANY_DOCS:
        if any(word in doc.lower() for word in query.split() if len(word) > 3):
            return {"context": doc}
    # If no match, return a random doc
    return {"context": random.choice(COMPANY_DOCS)}

@app.post("/voice")
def voice_endpoint(req: ChatRequest):
    """Voice endpoint for processing user messages (non-streaming, returns full JSON)."""
    try:
        from together import Together
    except ImportError:
        return {"response": "[Error: together SDK not installed on backend. Please install with: pip install together]"}
    TOGETHER_API_KEY = os.getenv("TOGETHER_API_KEY")
    if not TOGETHER_API_KEY:
        return {"response": "[Error: TOGETHER_API_KEY not set in environment]"}
    client = Together(api_key=TOGETHER_API_KEY)
    feedback_schema = {
        "type": "object",
        "properties": {
            "conversation": {"type": "string", "description": "Conversational message for the user."},
            "data": {
                "type": "object",
                "properties": {
                    "Q1": {"type": "object", "properties": {"answer": {"type": "string"}, "remark": {"type": "string"}}, "required": ["answer", "remark"]},
                    "Q2": {"type": "object", "properties": {"answer": {"type": "string"}, "remark": {"type": "string"}}, "required": ["answer", "remark"]},
                    "Q3": {"type": "object", "properties": {"answer": {"type": "string"}, "remark": {"type": "string"}}, "required": ["answer", "remark"]},
                    "Q4": {"type": "object", "properties": {"answer": {"type": "string"}, "remark": {"type": "string"}}, "required": ["answer", "remark"]},
                    "Q5": {"type": "object", "properties": {"answer": {"type": "string"}, "remark": {"type": "string"}}, "required": ["answer", "remark"]}
                },
                "required": ["Q1", "Q2", "Q3", "Q4", "Q5"]
            },
            "current_question": {"type": "string"},
            "is_complete": {"type": "boolean"}
        },
        "required": ["conversation", "data", "current_question", "is_complete"]
    }
    try:
        response = client.chat.completions.create(
            model="meta-llama/Llama-3.3-70B-Instruct-Turbo-Free",
            messages=req.messages,
            response_format={
                "type": "json_object",
                "schema": feedback_schema
            }
        )
        message_content = response.choices[0].message.content
        return {"response": message_content}
    except Exception as e:
        return {"response": f"[Error from Together API: {str(e)}]"}
